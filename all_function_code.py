"""사고관리 엑셀 → 사고사례 PPT 자동 생성: 전체 기능을 한 파일에 기능별 함수로 정리한 최신본.

실행:   streamlit run all_function_code.py          (Streamlit 화면. `streamlit run app.py`도 같은 화면)
필요 파일(같은 폴더): masked_사고관리시스템_이미지포함_사고내용보완.xlsx(없으면 화면에서 업로드), (ref1)AIglue_안전사고사례_PPT마스터슬라이드.pptx(아이콘 출처),
                     .env 또는 Streamlit Secrets(UPSTAGE_API_KEY=...; 없으면 조회만 되고 LLM 추출은 꺼진다)
설치:   pip install pandas openpyxl python-pptx pillow lxml requests streamlit

기능 구성(이 순서대로 아래에 나온다)
  1. 텍스트 폭·넘침 검사        text_units, wrapped_line_count, title_issue, body_issue ...
  2. 엑셀 → 표준 사고 레코드     build_standard_records, load_row_images ...
  3. 조회·통계·필터              missing_stats, unique_table, filter_cases, diagnose_empty ...
  4. 장소 규칙 추출              rule_based_place, places_differ
  5. Upstage API 클라이언트      load_api_key, chat_json
  6. LLM 추출 결과 캐시          open_cache, cache_get, cache_put
  7. LLM 추출(제목·장소·내용)     extract_case, quality_problems
  8. 슬라이드 값 매핑            default_team, level_code, default_place, to_case_slide
  9. PPT 생성                    build_presentation
 10. 조회·추출·PPT 로직          load_dataset, build_meta, run_query, extract_rows, build_ppt
 11. Streamlit 화면              render_streamlit_app (조회 → LLM 추출 → 확인·수정 표 → PPT)
"""
from __future__ import annotations

import copy
import datetime
import hashlib
import io
import json
import logging
import os
import re
import sys
import threading
import time
import webbrowser
import zipfile
from collections.abc import Callable, Mapping, Sequence
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

import pandas as pd
import requests
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

try:                                    # Streamlit 화면(11번)에서만 필요하다.
    import streamlit as st
except ImportError:
    st = None

log = logging.getLogger('accident_ppt')

# ====================================================================================================
# 0. 파일 위치·공통 상수
# ====================================================================================================
BASE = Path(__file__).resolve().parent
XLSX = BASE / 'masked_사고관리시스템_이미지포함_사고내용보완.xlsx'
TEMPLATE = BASE / '(ref1)AIglue_안전사고사례_PPT마스터슬라이드.pptx'      # PPT 아이콘 자산 출처
DOTENV = BASE / '.env'
EXTRACTION_CACHE = BASE / 'extraction_cache.json'

WEEKDAYS = '월화수목금토일'
MISSING_LABEL = '(미기재)'      # 엑셀 분류형 열의 결측 표준값
MISSING = '확인 필요'           # PPT 칸에 채울 값이 엑셀에 없을 때
UNKNOWN_PLACE = '파악불가'      # 사고내용에서 장소를 찾지 못했을 때
PLACE_JOINER = ' / '
PPTX_MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

# ====================================================================================================
# 1. 텍스트 폭·넘침 검사 (PPT 칸 크기·글꼴을 바꾸지 않고 글이 칸 안에 들어가는지 미리 확인)
#    한도는 PowerPoint에서 전각 한글 기준으로 실측한 값이며 Streamlit 화면의 검사도 같은 규칙을 쓴다.
# ====================================================================================================
# 전각(한글·한자·전각기호 등)은 1.0, 그 밖의 반각 글자(영문·숫자·공백·기호)는 0.6 단위로 센다
WIDE_CHARS = re.compile('[ᄀ-ᇿ⺀-꓏가-힣豈-﫿＀-｠￠-￦…·•]')
HALF_WIDTH_UNITS = 0.6
LINE_SAFETY_MARGIN = 0.5      # 줄 끝 단어 넘김 등으로 생기는 오차 여유
TITLE_PREFIX = '[사고사례] '


@dataclass(frozen=True)
class FieldLimits:
    title: float = 25.0        # 제목 25pt, 1줄 (접두 '[사고사례] ' 포함)
    date: float = 22.0
    place: float = 22.0
    team: float = 22.0
    level: float = 24.0
    body_line: float = 22.0    # 내용 칸 1줄
    body_max_lines: int = 6    # 내용 칸 최대 줄 수


LIMITS = FieldLimits()


def text_units(text: str) -> float:
    return sum(1.0 if WIDE_CHARS.match(ch) else HALF_WIDTH_UNITS for ch in text)


def wrapped_line_count(text: str, capacity: float) -> int:
    """글자 단위로 줄바꿈했을 때의 줄 수 (빈 문자열도 1줄)"""
    usable = capacity - LINE_SAFETY_MARGIN
    lines, used = 1, 0.0
    for ch in text:
        width = 1.0 if WIDE_CHARS.match(ch) else HALF_WIDTH_UNITS
        if used + width > usable:
            lines, used = lines + 1, 0.0
        used += width
    return lines


def bullet_line_count(bullets: Sequence[str], limits: FieldLimits = LIMITS) -> int:
    """'- ' 접두를 포함해 항목별로 줄바꿈한 총 줄 수"""
    return sum(wrapped_line_count('- ' + b, limits.body_line) for b in bullets)


def single_line_issue(label: str, text: str, limit: float) -> str | None:
    units = text_units(text)
    return f'{label}이(가) 너무 깁니다 ({units:.1f}/{limit:.0f})' if units > limit else None


def title_issue(title_words: str, limits: FieldLimits = LIMITS) -> str | None:
    """제목은 화면에 '[사고사례] ' + 단어로 표시되므로 접두를 포함해 검사한다."""
    return single_line_issue('제목', TITLE_PREFIX + title_words, limits.title)


def body_issue(bullets: Sequence[str], limits: FieldLimits = LIMITS) -> str | None:
    lines = bullet_line_count(bullets, limits)
    return f'내용이 {lines}줄로 한도({limits.body_max_lines}줄)를 넘습니다' if lines > limits.body_max_lines else None


# ====================================================================================================
# 2. 엑셀 → 표준 사고 레코드 (2단 헤더 처리 · 타입 정리 · 파생 컬럼 · 결측 표준화 · 떠 있는 이미지 추출)
# ====================================================================================================
SHEET = '사고관리시스템'
HEADER_ROWS = 2                                    # 엑셀 1~2행이 2단 헤더
FIRST_DATA_SHEET_ROW = HEADER_ROWS + 1
FILTER_VARS = ('사고유형', '재해정도', '사고성여부')  # 화면에서 통계·선택 대상으로 쓰는 변수
DateRange = tuple[datetime.date, datetime.date]

# 멀티인덱스 헤더(상단, 하단) → 표준 컬럼명. 줄바꿈·공백을 지운 뒤 비교한다.
HEADER_TO_STD = {
    ('No', None): 'No', ('일시', '날짜'): '날짜', ('일시', '시간'): '시간',
    ('소속', '부문'): '부문', ('소속', '현재팀'): '현재팀', ('소속', '등록당시팀'): '등록당시팀',
    ('소속', '책임부서'): '책임부서', ('소속', '관리부서'): '관리부서', ('프로젝트', None): '프로젝트',
    ('12대안전수칙', None): '안전수칙', ('사고내용', None): '사고내용', ('사고유형', None): '사고유형',
    ('재해정도', None): '재해정도', ('잠재등급', None): '잠재등급', ('실제등급', None): '실제등급',
    ('사고성여부', None): '사고성여부', ('조사단계', None): '조사단계', ('대책단계', None): '대책단계',
    ('이미지', None): '이미지',
}
TEXT_COLS = ['사고내용']
LABEL_COLS = ['부문', '현재팀', '등록당시팀', '책임부서', '관리부서', '프로젝트', '안전수칙',
              '사고유형', '재해정도', '잠재등급', '실제등급', '사고성여부', '조사단계', '대책단계']


def _clean_header(part) -> str | None:
    text = re.sub(r'\s+', '', str(part))
    return None if text.startswith('Unnamed') else text


def flatten_header(columns) -> list[str]:
    """2단 MultiIndex 헤더를 표준 컬럼명 리스트로 변환한다. 알 수 없는 헤더는 즉시 오류."""
    names = []
    for top, sub in columns:
        key = (_clean_header(top), _clean_header(sub))
        if key not in HEADER_TO_STD:
            raise KeyError(f'알 수 없는 헤더: {key} — 엑셀 양식이 바뀌었는지 확인하세요')
        names.append(HEADER_TO_STD[key])
    return names


def _clean_text(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return pd.NA
    text = str(value).strip()
    return text if text else pd.NA


def _normalize_time(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return pd.NA
    parsed = pd.to_datetime(str(value).strip(), errors='coerce')
    return pd.NA if pd.isna(parsed) else f'{parsed.hour:02d}:{parsed.minute:02d}'


ExcelSource = Path | bytes         # 로컬은 파일 경로, 클라우드는 사용자가 올린 파일의 bytes
MAX_ZIP_MEMBERS = 2000             # 공개 앱에 올라오는 파일 방어: 압축을 풀었을 때의 규모 상한 (실제 파일은 이보다 훨씬 작다)
MAX_ZIP_MEMBER_BYTES = 50 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 200 * 1024 * 1024
MAX_DATA_ROWS = 20_000


class ExcelFormatError(ValueError):
    """올린 파일이 이 앱이 기대하는 사고관리 엑셀 형식이 아님 (사용자에게 그대로 보여줄 수 있는 메시지)"""


def _excel_file(source: ExcelSource) -> Path | io.BytesIO:
    """경로는 그대로, 업로드된 bytes는 메모리 파일로 바꾼다 (pandas·zipfile 모두 파일 객체를 받는다)."""
    return io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source


def _check_excel_limits(data: bytes) -> None:
    """압축을 풀었을 때 비정상적으로 큰 파일(zip 폭탄)을 내용을 읽기 전에 거른다."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = z.infolist()
    if (len(members) > MAX_ZIP_MEMBERS or any(m.file_size > MAX_ZIP_MEMBER_BYTES for m in members)
            or sum(m.file_size for m in members) > MAX_ZIP_TOTAL_BYTES):
        raise ExcelFormatError('엑셀 파일의 구조나 크기가 정상 범위를 넘습니다. 사고관리 엑셀(.xlsx)을 올려 주세요.')


def _image_sheet_rows(xlsx: ExcelSource) -> set[int]:
    """이미지는 셀 값이 아니라 시트 위에 떠 있는 그림이므로 앵커 행(시트 행 번호)으로 찾는다."""
    with zipfile.ZipFile(_excel_file(xlsx)) as z:
        drawing = z.read('xl/drawings/drawing1.xml').decode('utf-8')
    return {int(row) + 1 for row in re.findall(r'<from><col>\d+</col>.*?<row>(\d+)</row>', drawing, flags=re.S)}


def load_row_images(xlsx: ExcelSource) -> dict[int, bytes]:
    """시트 행 번호 → 그 행의 이미지 bytes"""
    with zipfile.ZipFile(_excel_file(xlsx)) as z:
        drawing = z.read('xl/drawings/drawing1.xml').decode('utf-8')
        rel_xml = z.read('xl/drawings/_rels/drawing1.xml.rels').decode('utf-8')
        rels = {re.search(r'Id="([^"]+)"', tag).group(1): re.search(r'Target="([^"]+)"', tag).group(1)
                for tag in re.findall(r'<Relationship [^>]*>', rel_xml)}
        anchors = re.findall(r'<from><col>\d+</col>.*?<row>(\d+)</row>.*?r:embed="(rId\d+)"', drawing, flags=re.S)
        return {int(row) + 1: z.read(rels[rid].lstrip('/')) for row, rid in anchors}


def build_standard_records(xlsx: ExcelSource) -> pd.DataFrame:
    """엑셀을 표준 사고 레코드로 변환한다 (헤더 처리 · 타입 정리 · 파생 컬럼 · 결측 표준화)."""
    df = pd.read_excel(_excel_file(xlsx), sheet_name=SHEET, header=list(range(HEADER_ROWS)), dtype=object)
    df.columns = flatten_header(df.columns)

    df['No'] = pd.to_numeric(df['No'], errors='coerce').astype('Int64')
    df['날짜'] = pd.to_datetime(df['날짜'], errors='coerce')
    df['시간'] = df['시간'].map(_normalize_time).astype('string')
    for col in TEXT_COLS + LABEL_COLS:
        df[col] = df[col].map(_clean_text).astype('string')

    df['시트행'] = range(FIRST_DATA_SHEET_ROW, FIRST_DATA_SHEET_ROW + len(df))
    df['일시'] = pd.to_datetime(df['날짜'].dt.strftime('%Y-%m-%d') + ' ' + df['시간'], errors='coerce').fillna(df['날짜'])
    df['요일'] = df['날짜'].dt.dayofweek.map(dict(enumerate(WEEKDAYS))).astype('string')
    df['주_시작일'] = df['날짜'].dt.to_period('W-SUN').dt.start_time
    df['연도'] = df['날짜'].dt.year.astype('Int64')
    df['이미지_유무'] = df['시트행'].isin(_image_sheet_rows(xlsx))

    df[TEXT_COLS] = df[TEXT_COLS].fillna('')
    df[LABEL_COLS] = df[LABEL_COLS].fillna(MISSING_LABEL)
    columns = (['No', '시트행', '날짜', '시간', '일시', '요일', '주_시작일', '연도'] + LABEL_COLS[:6]
               + ['안전수칙', '사고내용'] + LABEL_COLS[7:] + ['이미지_유무'])
    return df[columns]


# ====================================================================================================
# 3. 조회·통계·필터 (일 단위 기간 + 변수별 선택, 결측 현황, 고유값 목록, 0건 원인 진단)
# ====================================================================================================
def missing_stats(df: pd.DataFrame, columns: Sequence[str] = FILTER_VARS) -> pd.DataFrame:
    """변수별 (전체 건수, 결측 건수, 결측 비율). 결측 = 표준화된 '(미기재)' 또는 빈 문자열."""
    total = len(df)
    rows = []
    for col in columns:
        missing = int(((df[col] == MISSING_LABEL) | (df[col] == '')).sum())
        rows.append({'변수': col, '전체 건수': total, '결측 건수': missing,
                     '결측 비율': missing / total if total else 0.0})
    return pd.DataFrame(rows)


def unique_table(df: pd.DataFrame, column: str) -> pd.DataFrame:
    """변수의 고유값 표 (값, 건수, 비율). 건수 많은 순, 동률은 가나다순, '(미기재)'는 맨 뒤."""
    counts = df[column].value_counts()
    ranked = sorted(counts.index, key=lambda v: (v == MISSING_LABEL, -counts[v], v))
    total = len(df)
    return pd.DataFrame({'값': ranked, '건수': [int(counts[v]) for v in ranked],
                         '비율': [counts[v] / total for v in ranked]})


def parse_date(value) -> datetime.date:
    """date · datetime · 'YYYY-MM-DD' 문자열을 날짜로 변환한다. 해석할 수 없으면 ValueError."""
    parsed = pd.Timestamp(value)
    if pd.isna(parsed):
        raise ValueError(f'날짜를 해석할 수 없습니다: {value!r}')
    return parsed.date()


def normalize_date_range(start, end) -> DateRange:
    """(시작일, 종료일) 검증. 하루만 고르려면 두 값을 같게 준다 (양 끝 포함)."""
    first, last = parse_date(start), parse_date(end)
    if first > last:
        raise ValueError(f'시작일({first})이 종료일({last})보다 늦습니다')
    return first, last


def date_bounds(df: pd.DataFrame) -> DateRange:
    """달력에 줄 선택 가능 범위 (첫 날, 마지막 날)"""
    return df['날짜'].min().date(), df['날짜'].max().date()


def cases_per_day(df: pd.DataFrame) -> dict[str, int]:
    """날짜(ISO 문자열) → 그날의 건수. 사례가 없는 날은 키가 없다."""
    counts = df.groupby(df['날짜'].dt.date).size()
    return {day.isoformat(): int(n) for day, n in counts.items()}


MATCH_MODES = ('AND', 'OR')          # 변수 사이의 결합 방식 (기본 AND)


def filter_cases(df: pd.DataFrame, selections: Mapping[str, Sequence[str]],
                 date_range: DateRange | None = None, match: str = 'AND') -> pd.DataFrame:
    """선택 조건에 맞는 케이스를 시트행 오름차순으로 반환한다 (df는 변경하지 않음).

    date_range: (시작일, 종료일) 일 단위, 양 끝 포함. None이면 날짜 조건 없음. 날짜는 항상 다른 조건과 AND.
    selections: {변수: [선택값, ...]}. 변수 안의 값은 OR. 비어 있는(=전체) 변수는 조건 없음으로 무시한다.
    match: 변수 사이의 결합. 'AND'(기본)=모두 만족, 'OR'=하나라도 만족. 선택된 변수가 없으면 결합할 조건이 없다.
    엑셀에 없는 값이 들어오면 조용히 0건이 되지 않도록 ValueError를 낸다.
    """
    if match not in MATCH_MODES:
        raise ValueError(f'match는 {MATCH_MODES} 중 하나여야 합니다: {match!r}')
    mask = pd.Series(True, index=df.index)
    if date_range is not None:
        start, end = normalize_date_range(*date_range)
        mask &= (df['날짜'] >= pd.Timestamp(start)) & (df['날짜'] < pd.Timestamp(end) + pd.Timedelta(days=1))
    checks = []
    for column, chosen in selections.items():
        chosen = list(chosen)
        if not chosen:
            continue
        known = set(df[column].unique())
        unknown = [v for v in chosen if v not in known]
        if unknown:
            raise ValueError(f'{column}에 없는 값: {unknown} (가능한 값: {sorted(known)})')
        checks.append(df[column].isin(chosen))
    if checks:
        combined = checks[0]
        for check in checks[1:]:
            combined = combined & check if match == 'AND' else combined | check
        mask &= combined
    return df[mask].sort_values('시트행').reset_index(drop=True)


def diagnose_empty(df: pd.DataFrame, selections: Mapping[str, Sequence[str]],
                   date_range: DateRange | None = None, match: str = 'AND') -> dict[str, int]:
    """0건일 때 원인 파악용: 조건을 하나씩 뺐을 때의 건수."""
    active = {col: list(vals) for col, vals in selections.items() if list(vals)}
    report = {}
    for column in active:
        rest = {c: v for c, v in active.items() if c != column}
        report[f'{column} 조건을 뺀 경우'] = len(filter_cases(df, rest, date_range, match))
    if date_range is not None:
        report['일자 조건을 뺀 경우'] = len(filter_cases(df, active, None, match))
    return report


# ====================================================================================================
# 4. 장소 규칙 추출 (LLM 추출값과 비교하는 기준값)
# ====================================================================================================
# 문장 첫머리의 "…에서" 앞부분이 장소 (예: '조립1공장 3번 정반에서 …', '옥외 자재적치장에서 …')
_PLACE_PATTERN = re.compile(r'^\s*(?P<place>[^,.\n]{2,40}?)(?:에서는|에서|에서의)\s')


def rule_based_place(text: str) -> str | None:
    """장소를 찾으면 그 문구, 못 찾거나 칸(1줄)에 안 들어가는 길이면 None."""
    match = _PLACE_PATTERN.match(text or '')
    if not match:
        return None
    place = match.group('place').strip()
    return place if place and text_units(place) <= LIMITS.place else None


def normalize_place(place: str | None) -> str:
    """비교용: 공백 제거, 없으면 빈 문자열"""
    return re.sub(r'\s+', '', place or '')


def places_differ(rule: str | None, llm: str) -> bool:
    """규칙 추출값이 있고 LLM 값과 다르면 True. 규칙이 못 찾은 경우는 비교하지 않는다."""
    return rule is not None and normalize_place(rule) != normalize_place(llm)


# ====================================================================================================
# 5. Upstage Solar API 클라이언트 (Chat Completions + JSON Schema 구조화 출력)
#    문서: https://console.upstage.ai/api/chat  (모델: solar-mini / solar-pro2 / solar-pro3 / solar-pro4)
#    API 키는 환경변수 UPSTAGE_API_KEY 또는 .env에서 읽고, 코드·로그·화면에는 절대 남기지 않는다.
# ====================================================================================================
API_URL = 'https://api.upstage.ai/v1/chat/completions'
KEY_ENV_NAME = 'UPSTAGE_API_KEY'
RETRY_STATUSES = {429, 500, 502, 503, 504}


class ApiKeyMissing(RuntimeError):
    """API 키가 설정되지 않음"""


class UpstageError(RuntimeError):
    def __init__(self, status: int, message: str) -> None:
        super().__init__(f'Upstage API 오류 {status}: {message}')
        self.status = status
        self.message = message


class ModelUnavailable(UpstageError):
    """요청한 모델을 쓸 수 없음(없음·권한 없음·구조화 출력 미지원 등) → 다른 모델로 대체할 수 있는 오류"""


@dataclass(frozen=True)
class ChatResult:
    content: dict
    model: str
    total_tokens: int


@dataclass(frozen=True)
class UpstageClient:
    api_key: str = field(repr=False)                          # repr에도 키가 나오지 않게 한다
    session: requests.Session = field(default_factory=requests.Session)
    timeout: float = 60.0
    max_retries: int = 5
    sleep: Callable[[float], None] = time.sleep


SECRET_PATTERNS = (re.compile(r'(?<![A-Za-z0-9])up_[A-Za-z0-9]{16,}'),          # '_'는 경계로 보지 않는다(접두어 뒤에 붙은 키도 잡는다)
                   re.compile(r'Bearer\s+[A-Za-z0-9._\-+/=]{8,}', re.IGNORECASE))
REDACTED = '***'


def redact_secrets(text: str, secrets: Sequence[str] = ()) -> str:
    """오류 메시지·로그에 API 키가 섞여 나오지 않게 가린다 (정확한 키 값 + 키 모양 토큰 + Bearer 토큰)."""
    for secret in secrets:
        if secret and secret.strip():
            text = text.replace(secret.strip(), REDACTED)
    for pattern in SECRET_PATTERNS:
        text = pattern.sub(REDACTED, text)
    return text


def _streamlit_secrets() -> dict[str, object]:
    """Streamlit Secrets(클라우드의 Secrets 창, 로컬의 .streamlit/secrets.toml)에서 키 항목만 꺼낸다.
    Secrets가 없거나 Streamlit이 없으면 빈 값이다. 다른 비밀 항목은 읽지 않는다."""
    if st is None:
        return {}
    try:
        return {KEY_ENV_NAME: st.secrets[KEY_ENV_NAME]} if KEY_ENV_NAME in st.secrets else {}
    except Exception as error:                       # Secrets 파일이 없을 때의 예외 종류가 버전마다 달라 넓게 받는다
        log.debug('Streamlit Secrets를 읽지 못했습니다: %s', type(error).__name__)
        return {}


def load_api_key(env: Mapping[str, str] | None = None, dotenv_path: Path | None = None,
                 secrets: Mapping[str, object] | None = None) -> str:
    """환경변수 → .env → Streamlit Secrets 순서로 키를 찾는다. 없으면 ApiKeyMissing."""
    env = os.environ if env is None else env
    key = env.get(KEY_ENV_NAME, '').strip()
    if not key and dotenv_path and dotenv_path.exists():
        for line in dotenv_path.read_text(encoding='utf-8-sig').splitlines():
            name, _, value = line.strip().partition('=')
            if name.strip() == KEY_ENV_NAME:
                key = value.strip().strip('"').strip("'")
    if not key and secrets:
        value = secrets.get(KEY_ENV_NAME)
        key = value.strip() if isinstance(value, str) else ''
    if not key:
        raise ApiKeyMissing(f'{KEY_ENV_NAME}가 설정되지 않았습니다. 환경변수, .env 파일, Streamlit Secrets 중 한 곳에 넣어 주세요.')
    return key


def _backoff(attempt: int, retry_after: str | None = None) -> float:
    """재시도 대기 시간(초): 서버가 Retry-After를 주면 그 값, 아니면 3·6·12·24·30초(상한 30초)"""
    try:
        if retry_after:
            return min(60.0, max(1.0, float(retry_after)))
    except ValueError:
        pass
    return min(30.0, 3.0 * 2 ** attempt)


def _error_message(response: requests.Response, secrets: Sequence[str] = ()) -> str:
    """서버가 돌려준 오류 문구. 키가 섞여 있을 수 있어 먼저 가리고 나서 300자로 자른다(경계에서 키 조각이 남지 않게)."""
    try:
        data = response.json()
        detail = data.get('error', data) if isinstance(data, dict) else data
        text = detail.get('message') if isinstance(detail, dict) else str(detail)
    except ValueError:
        text = response.text
    return redact_secrets(str(text or response.reason or '알 수 없는 오류'), secrets)[:300]


def _mentions_model_or_format(message: str, body: dict) -> bool:
    lowered = message.lower()
    return any(word in lowered for word in ('model', 'response_format', 'json_schema', 'schema', 'structured', body['model']))


def _post_with_retry(client: UpstageClient, body: dict) -> dict:
    """429·5xx·네트워크 오류는 대기 후 재시도, 모델/형식 문제는 ModelUnavailable, 그 밖은 UpstageError"""
    headers = {'Authorization': f'Bearer {client.api_key}', 'Content-Type': 'application/json'}
    for attempt in range(client.max_retries + 1):
        try:
            response = client.session.post(API_URL, headers=headers, json=body, timeout=client.timeout)
        except requests.RequestException as error:
            if attempt == client.max_retries:
                raise UpstageError(0, f'네트워크 오류: {type(error).__name__}') from error
            client.sleep(_backoff(attempt))
            continue
        if response.status_code == 200:
            return response.json()
        message = _error_message(response, (client.api_key,))
        if response.status_code in RETRY_STATUSES and attempt < client.max_retries:
            log.debug('Upstage 일시 오류 %s, 재시도 %d/%d', response.status_code, attempt + 1, client.max_retries)
            client.sleep(_backoff(attempt, response.headers.get('Retry-After')))
            continue
        if response.status_code in (400, 401, 403, 404, 422) and _mentions_model_or_format(message, body):
            raise ModelUnavailable(response.status_code, message)
        raise UpstageError(response.status_code, message)
    raise UpstageError(0, '재시도 횟수를 초과했습니다.')


def chat_json(client: UpstageClient, model: str, messages: list[dict], schema_name: str, schema: dict, *,
              temperature: float = 0.0, max_tokens: int = 600) -> ChatResult:
    """JSON Schema(strict)를 지정해 호출하고, 스키마를 따른 JSON(dict)을 돌려준다."""
    body = {
        'model': model, 'messages': messages, 'temperature': temperature, 'max_tokens': max_tokens,
        'response_format': {'type': 'json_schema',
                            'json_schema': {'name': schema_name, 'strict': True, 'schema': schema}},
    }
    payload = _post_with_retry(client, body)
    try:
        content = json.loads(payload['choices'][0]['message']['content'])
        tokens = int((payload.get('usage') or {}).get('total_tokens') or 0)
    except (KeyError, IndexError, TypeError, ValueError) as error:
        raise UpstageError(502, '응답 형식이 올바르지 않습니다.') from error
    if not isinstance(content, dict):
        raise UpstageError(502, '응답이 JSON 객체가 아닙니다.')
    return ChatResult(content=content, model=payload.get('model', model), total_tokens=tokens)


# ====================================================================================================
# 6. LLM 추출 결과 캐시 (같은 사고내용을 다시 조회할 때 API를 다시 부르지 않기 위함)
# ====================================================================================================
@dataclass
class ExtractionCache:
    path: Path
    data: dict[str, dict] = field(default_factory=dict)
    lock: threading.Lock = field(default_factory=threading.Lock, repr=False)


def open_cache(path: Path) -> ExtractionCache:
    """캐시 파일을 읽는다. 없거나 손상됐으면 빈 캐시로 시작한다(앱 전체가 멈추지 않게)."""
    data: dict = {}
    if path.exists():
        try:
            loaded = json.loads(path.read_text(encoding='utf-8'))
            data = loaded if isinstance(loaded, dict) else {}
        except (OSError, ValueError) as error:
            log.warning('추출 캐시를 읽지 못해 빈 캐시로 시작합니다: %s', type(error).__name__)
    return ExtractionCache(path=path, data=data)


def cache_key(prompt_version: str, text: str) -> str:
    return f'{prompt_version}:{hashlib.sha1(text.encode("utf-8")).hexdigest()[:16]}'


def cache_get(cache: ExtractionCache, key: str) -> dict | None:
    with cache.lock:
        entry = cache.data.get(key)
        return dict(entry) if entry else None


def cache_put(cache: ExtractionCache, key: str, entry: dict) -> None:
    with cache.lock:
        cache.data[key] = dict(entry)
        temp = cache.path.with_suffix('.tmp')
        temp.write_text(json.dumps(cache.data, ensure_ascii=False, indent=1), encoding='utf-8')
        os.replace(temp, cache.path)                 # 쓰다 끊겨도 기존 파일이 깨지지 않도록 교체 방식으로 저장


# ====================================================================================================
# 7. LLM 추출: 사고내용 → 제목·장소·내용(개조식)
#    solar-mini를 먼저 쓰고, ① mini를 쓸 수 없거나 ② 품질 검사를 재시도 후에도 통과하지 못하면 solar-pro4로 올린다.
#    그래도 남는 문제는 warnings로 돌려주어 화면에서 사용자가 고치게 한다.
# ====================================================================================================
PROMPT_VERSION = 'v3'
PRIMARY_MODEL = 'solar-mini'
FALLBACK_MODEL = 'solar-pro4'
ATTEMPTS_PER_MODEL = 2
MIN_BULLETS, MAX_BULLETS = 2, 3

# 필요한 항목은 모두 required (strict 모드: additionalProperties=false)
SCHEMA = {
    'type': 'object',
    'properties': {
        'title': {'type': 'string', 'description': '사고를 압축한 단어 위주 제목'},
        'place': {'type': 'string', 'description': '사고 장소. 알 수 없으면 정확히 "파악불가"'},
        'bullets': {'type': 'array', 'items': {'type': 'string'}, 'description': '개조식 요약 2~3개'},
    },
    'required': ['title', 'place', 'bullets'],
    'additionalProperties': False,
}

SYSTEM_PROMPT = f"""당신은 제조·조선 현장 안전사고 보고서를 교육용 슬라이드 문구로 정리하는 편집자입니다.
"사고내용" 원문을 읽고 JSON(title, place, bullets)으로만 답하세요.

[title: 제목]
- 사고를 압축적으로 전달하는 단어 위주의 명사구입니다. 문장·서술어 금지.
- 예: "프레스 금형 교체 중 손 협착", "그라인더 작업 중 손가락 베임"
- 공백 포함 15자 이내. '사고' 같은 군더더기 단어는 쓰지 않습니다.

[place: 장소]
- 사고가 발생한 장소(공장·구역·설비·위치)를 원문의 표현 그대로 씁니다. 공백 포함 20자 이내.
- 원문에서 알 수 없으면 정확히 "{UNKNOWN_PLACE}"라고 씁니다. 추측하지 않습니다.

[bullets: 내용]
- 개조식 문장 {MIN_BULLETS}~{MAX_BULLETS}개. 원문이 짧으면 {MIN_BULLETS}개, 길면 {MAX_BULLETS}개.
- 개조식이란 서술어("~했다", "~하였다", "~되었다", "~입었다")를 없애고 명사·명사구로 끝내는 문체입니다.
  예) "화상이 발생했다" → "화상 발생", "부상을 입었다" → "부상"
- 핵심 단어 중심으로 쓰되 원문에 충실해야 합니다.
  · 원문의 주요 내용을 누락하지 않습니다. 원문에 있는 다음 요소는 반드시 포함합니다.
    ① 어떤 작업·상황에서 ② 무슨 일이 일어났는지(경위·원인) ③ 결과(부상 부위·증상·피해)
    ④ 조치·현재 상태(병원 이송, 응급 처치, 생명 지장 없음, 추가 부상 없음 등)
  · 원문에 없는 내용은 절대 추가하지 않습니다.
  · 호기·번호·신체 부위·수치(예: 3호기, 2미터) 등 구체적 표현은 그대로 유지합니다.
- 각 항목은 서로 다른 사실을 담은 하나의 개조식 구절입니다. 장소만 적거나 조사(~에서, ~하며)로 끝나는
  문장 조각은 금지합니다. '사고 발생', '사고를 당함'처럼 내용이 없는 표현도 쓰지 않습니다.
- '&', '→' 같은 기호는 쓰지 않고 쉼표를 사용합니다.
- 항목 전체 글자 수 합계는 공백 포함 약 110자 이내로 합니다.

[개조식 변환 예시]
예시 1
- 원문: 취부 작업 중 바닥에 떨어진 절단 불티를 인지하지 못하고 손바닥으로 짚은 순간 화상이 발생했다
- 개조식: 취부 작업 중 바닥에 떨어진 절단 불티를 인지하지 못하고 손바닥으로 짚은 순간 화상 발생
예시 2
- 원문: 좌측 두 번째 수지(손가락)에 화상 부상을 입었다
- 개조식: 좌측 두 번째 수지(손가락) 화상 부상"""

# 출력 형식·문체를 보여주는 대화 예시 (장소가 원문에 없으면 파악불가로 쓰는 경우 포함)
FEW_SHOT = [
    {'role': 'user', 'content': '사고내용 원문:\n프레스 금형 교체작업 중 안전핀 미체결 상태에서 페달이 오작동되어 작업자 우측 손가락이 '
                              '협착되었다. 병원으로 이송되어 응급 수술을 받았으며 3지골 골절로 4주 진단을 받았다.'},
    {'role': 'assistant', 'content': json.dumps({
        'title': '프레스 금형 교체 중 손 협착', 'place': UNKNOWN_PLACE,
        'bullets': ['금형 교체작업 중 안전핀 미체결 상태에서 페달 오작동으로 작업자 우측 손가락 협착',
                    '병원 이송 후 응급 수술 시행', '3지골 골절, 4주 진단']}, ensure_ascii=False)},
]

# 원문에 이 말이 있으면 결과에도 (같은 말 또는 별칭이) 있어야 한다: 부상·조치·상태 같은 핵심 내용 누락 방지
CRITICAL_TERMS = {'생명': ('생명', '지장'), '병원': ('병원', '이송'), '이송': ('이송', '병원', '후송'), '응급': ('응급',),
                  '수술': ('수술',), '골절': ('골절',), '입원': ('입원',), '사망': ('사망',), '통증': ('통증',),
                  '어지러움': ('어지러움', '어지럼'), '출혈': ('출혈',), '화상': ('화상',)}
FRAGMENT_ENDINGS = ('에서', '에', '으로', '로', '하며', '하고', '며', '고', '서', '와', '과', '및', ',')
MIN_FRAGMENT_UNITS = 14        # '~ 작업 중'처럼 짧게 끝나는 항목은 사건이 빠진 조각으로 본다


@dataclass(frozen=True)
class Extraction:
    title: str                     # 접두 '[사고사례] '를 뺀 제목 단어
    place_llm: str                 # 알 수 없으면 '파악불가'
    bullets: tuple[str, ...]       # '- ' 접두 없는 개조식 항목
    model: str
    tokens: int = 0                # 이 추출에 쓴 총 토큰 (재시도 포함, 캐시에서 읽은 경우 저장된 값)
    warnings: tuple[str, ...] = field(default=())    # 재시도 후에도 남은 문제 → 화면에서 사용자가 수정


@dataclass
class CaseExtractor:
    """추출기 상태: API 클라이언트, (선택) 캐시, 사용 모델. mini를 쓸 수 없다고 판명되면 이후 pro만 쓴다."""
    client: UpstageClient
    cache: ExtractionCache | None = None
    primary_model: str = PRIMARY_MODEL
    fallback_model: str = FALLBACK_MODEL
    primary_unavailable: bool = False


def _clean_bullet(text: str) -> str:
    return text.strip().lstrip('-·•').strip()


def _normalize_llm_output(raw: dict) -> tuple[str, str, list[str]]:
    title = str(raw.get('title', '')).strip()
    if title.startswith(TITLE_PREFIX.strip()):
        title = title[len(TITLE_PREFIX.strip()):].strip()
    place = str(raw.get('place', '')).strip() or UNKNOWN_PLACE
    bullets = [b for b in (_clean_bullet(str(item)) for item in raw.get('bullets') or []) if b]
    return title, place, bullets


def quality_problems(source: str, title: str, place: str, bullets: list[str]) -> tuple[list[str], list[str]]:
    """(품질 문제 — 재시도·상위 모델 필요, 길이 문제 — 사용자가 고칠 수 있음)"""
    hard, soft = [], []
    if not title:
        hard.append('title이 비어 있습니다')
    elif (issue := title_issue(title)):
        soft.append(f'title {issue}: 15자 이내의 단어 위주로 줄이세요')
    if (issue := single_line_issue('place', place, LIMITS.place)):
        soft.append(f'{issue}: 20자 이내로 줄이세요')
    if not MIN_BULLETS <= len(bullets) <= MAX_BULLETS:
        hard.append(f'bullets는 {MIN_BULLETS}~{MAX_BULLETS}개여야 합니다 (현재 {len(bullets)}개)')
        return hard, soft
    if (issue := body_issue(bullets)):
        hard.append(f'{issue}: 항목을 더 압축하세요 (합계 약 110자 이내)')
    place_key = normalize_place(place)
    for bullet in bullets:
        if bullet.endswith(FRAGMENT_ENDINGS) or (bullet.endswith('중') and text_units(bullet) < MIN_FRAGMENT_UNITS) \
                or (place_key and normalize_place(bullet) == place_key):
            hard.append(f'"{bullet}"는 사건·결과가 없는 문장 조각입니다. 완결된 개조식 구절로 쓰세요')
        if re.search(r'[&→]', bullet):
            hard.append(f'"{bullet}"에서 기호(&, →) 대신 쉼표를 쓰세요')
    result_text = normalize_place(title + place + ''.join(bullets))
    result_numbers = set(re.findall(r'\d+', result_text))
    missing_numbers = sorted({n for n in re.findall(r'\d+', source) if n not in result_numbers})
    if missing_numbers:
        hard.append(f'원문의 숫자 {", ".join(missing_numbers)}가 빠졌습니다. 호기·번호·수치를 유지하세요')
    missing_terms = [term for term, aliases in CRITICAL_TERMS.items()
                     if term in source and not any(alias in result_text for alias in aliases)]
    if missing_terms:
        hard.append(f'원문의 핵심 내용({", ".join(missing_terms)})이 누락되었습니다')
    return hard, soft


def _call_models(extractor: CaseExtractor, text: str) -> Extraction:
    """mini → (재시도) → pro4 순서로 호출한다. 문제가 남으면 마지막 결과와 문제 목록을 warnings로 돌려준다."""
    tokens = 0
    best: tuple[str, str, list[str], str] | None = None
    problems: list[str] = []
    models = [extractor.fallback_model] if extractor.primary_unavailable \
        else [extractor.primary_model, extractor.fallback_model]
    for stage, model in enumerate(models):
        messages = [{'role': 'system', 'content': SYSTEM_PROMPT}, *FEW_SHOT,
                    {'role': 'user', 'content': f'사고내용 원문:\n{text}'}]
        for attempt in range(ATTEMPTS_PER_MODEL):
            try:
                chat = chat_json(extractor.client, model, messages, 'accident_case', SCHEMA)
            except ModelUnavailable as error:
                if stage == len(models) - 1:
                    raise
                log.warning('%s 사용 불가(%s) → %s로 전환합니다.', model, error.message, extractor.fallback_model)
                extractor.primary_unavailable = True
                break                                   # 다음 단계(상위 모델)로
            tokens += chat.total_tokens
            title, place, bullets = _normalize_llm_output(chat.content)
            hard, soft = quality_problems(text, title, place, bullets)
            best, problems = (title, place, bullets, chat.model), hard + soft
            if not hard and (not soft or attempt == ATTEMPTS_PER_MODEL - 1):
                return Extraction(title, place, tuple(bullets), chat.model, tokens, tuple(soft))
            log.info('추출 결과 재요청(%s %d/%d): %s', model, attempt + 1, ATTEMPTS_PER_MODEL, problems)
            messages += [{'role': 'assistant', 'content': json.dumps(chat.content, ensure_ascii=False)},
                         {'role': 'user', 'content': '다음 문제를 고쳐서 같은 JSON 형식으로 다시 답하세요: ' + '; '.join(problems)}]
    title, place, bullets, model = best
    return Extraction(title, place, tuple(bullets), model, tokens, tuple(problems))


def extract_case(extractor: CaseExtractor, text: str) -> Extraction:
    """사고내용 1건 → 제목·장소·내용. 캐시에 있으면 API를 부르지 않는다(품질 검사 경고는 다시 계산)."""
    key = cache_key(PROMPT_VERSION, text)
    if extractor.cache and (hit := cache_get(extractor.cache, key)):
        title, place, bullets = hit['title'], hit['place_llm'], list(hit['bullets'])
        hard, soft = quality_problems(text, title, place, bullets)
        return Extraction(title, place, tuple(bullets), hit['model'], hit.get('tokens', 0), tuple(hard + soft))
    result = _call_models(extractor, text)
    if extractor.cache:
        cache_put(extractor.cache, key, {'title': result.title, 'place_llm': result.place_llm,
                                         'bullets': list(result.bullets), 'model': result.model, 'tokens': result.tokens})
    return result


def build_extractor() -> CaseExtractor | None:
    """API 키가 있으면 Upstage 추출기, 없으면 None (화면은 LLM 없이 조회만 가능하다고 안내)"""
    try:
        key = load_api_key(dotenv_path=DOTENV, secrets=_streamlit_secrets())
    except ApiKeyMissing as error:
        log.warning('%s', error)
        return None
    return CaseExtractor(UpstageClient(key), open_cache(EXTRACTION_CACHE))


# ====================================================================================================
# 8. 슬라이드 값 매핑 (칸별 규칙을 한곳에)
#    제목='[사고사례] '+LLM 제목(수정 가능) / 일시=엑셀 날짜·시간 / 장소=규칙·LLM(수정 가능, 없으면 '파악불가')
#    소속=엑셀 관리부서 / 책임부서 그대로(협력사 포함) / 내용=LLM 개조식 / 재해정도=엑셀 재해정도 코드 그대로 / 사진=엑셀 이미지
# ====================================================================================================
@dataclass(frozen=True)
class CaseSlide:
    """사고사례 슬라이드 1장에 들어갈 값"""
    title: str
    date: str
    place: str
    team: str
    bullets: tuple[str, ...]
    level: str
    photo: bytes | None = None


def _present(value) -> str:
    text = '' if value is None else str(value).strip()
    return '' if text in ('', MISSING_LABEL) else text


def format_datetime(date_value, time_value) -> str:
    day = datetime.date.fromisoformat(str(date_value)[:10])
    return f'{day:%Y.%m.%d}({WEEKDAYS[day.weekday()]}) {str(time_value)[:5]}'


def default_team(record: dict) -> str:
    """관리부서 / 책임부서를 적힌 그대로 (둘이 같으면 한 번만, 둘 다 없으면 '확인 필요')"""
    parts: list[str] = []
    for column in ('관리부서', '책임부서'):
        value = _present(record.get(column))
        if value and value not in parts:
            parts.append(value)
    return PLACE_JOINER.join(parts) or MISSING


def level_code(record: dict) -> str:
    """재해정도 코드 원문 그대로 (없으면 '확인 필요')"""
    return _present(record.get('재해정도')) or MISSING


def default_place(rule: str | None, llm: str) -> tuple[str, bool]:
    """(칸에 처음 채울 값, 규칙과 LLM이 달라 사용자 승인이 필요한지). 다르면 두 값을 모두 적는다: '규칙값 / LLM값'."""
    if rule is None or normalize_place(rule) == normalize_place(llm):
        return llm, False
    return f'{rule}{PLACE_JOINER}{llm}', True


def title_text(title_words: str) -> str:
    return TITLE_PREFIX + title_words.strip()


def to_case_slide(record: dict, *, title: str, place: str, team: str, bullets: Sequence[str],
                  photo: bytes | None) -> CaseSlide:
    cleaned = tuple('- ' + b.strip().lstrip('-·•').strip() for b in bullets if b.strip())
    return CaseSlide(title=title_text(title), date=format_datetime(record['날짜'], record['시간']),
                     place=place.strip(), team=team.strip(), bullets=cleaned,
                     level=f'재해정도 : {level_code(record)}', photo=photo)


# ====================================================================================================
# 9. PPT 생성: 템플릿(표지·목차·구분·사고사례)을 python-pptx로 코드에서 그린다.
#    아이콘 PNG 6개만 템플릿 파일에서 꺼내 쓰고, 나머지 도형·텍스트·색·위치는 모두 코드로 재현한다.
# ====================================================================================================
DARK, YELLOW, WHITE = '1C1C1C', 'FDB913', 'FFFFFF'
GRAY_LIGHT, GRAY_MID = 'A8A8A8', '6E6E6E'
PANEL, RULE, RED, PHOTO_BG = 'EDEDEA', 'C9C9C9', 'D32F2F', 'F7F7F7'
SLIDE_W, SLIDE_H = 9144000, 6858000            # 25.4 x 19.05cm (4:3)
ICON_FILES = {'warning': 'image1.png', 'calendar': 'image2.png', 'pin': 'image3.png',
              'building': 'image4.png', 'clipboard': 'image5.png', 'camera': 'image6.png'}
CASE_ROWS = [('calendar', '일시', 1920240), ('pin', '장소', 2331720),
             ('building', '소속', 2743200), ('clipboard', '내용', 3154680)]   # (아이콘, 라벨, y)
PHOTO_FRAME = (5486400, 1691640, 3200400, 3429000)                            # 사진 영역 x, y, w, h
SECTION_ITEM = ('01', '당사 주간 사고사례', '이번 주 발생한 사내/협력사 안전사고 개요')
_ALIGN = {'l': PP_ALIGN.LEFT, 'ctr': PP_ALIGN.CENTER}
_ANCHOR = {'t': MSO_ANCHOR.TOP, 'ctr': MSO_ANCHOR.MIDDLE}


@lru_cache(maxsize=1)
def _icons(template: Path) -> dict[str, bytes]:
    if not template.exists():
        raise FileNotFoundError(f'아이콘을 꺼낼 템플릿 파일이 없습니다: {template}')
    with zipfile.ZipFile(template) as z:
        return {name: z.read(f'ppt/media/{file}') for name, file in ICON_FILES.items()}


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _new_slide(prs, bg: str = WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(bg)
    return slide


def _add_rect(slide, x, y, w, h, fill, line=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    style = shp._element.find(qn('p:style'))
    if style is not None:
        shp._element.remove(style)
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    if line:
        color, width, dashed = line
        shp.line.color.rgb, shp.line.width = _rgb(color), Emu(width)
        if dashed:
            shp.line.dash_style = MSO_LINE.DASH
    else:
        shp.line.fill.background()
    if shadow:
        etree.SubElement(shp._element.spPr, qn('a:effectLst')).append(etree.fromstring(
            '<a:outerShdw xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'blurRad="76200" dist="25400" dir="5400000" algn="bl" rotWithShape="0">'
            '<a:srgbClr val="808080"><a:alpha val="30000"/></a:srgbClr></a:outerShdw>'))
    return shp


def _para(text, size, color, bold=False, italic=False, align='l', line_pt=None, spc=None, kern=None) -> dict:
    return dict(text=text, size=size, color=color, bold=bold, italic=italic, align=align,
                line_pt=line_pt, spc=spc, kern=kern)


def _add_text(slide, x, y, w, h, paragraphs, anchor='t', margins=(0, 0, 0, 0), name=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE              # 글이 길어도 상자 크기·레이아웃을 바꾸지 않는다 (넘침은 생성 전에 검사)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Emu(m) for m in margins)
    tf.vertical_anchor = _ANCHOR[anchor]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[para['align']]
        if para['line_pt']:
            p.line_spacing = Pt(para['line_pt'])
        run = p.add_run()
        run.text = para['text']
        font = run.font
        font.size, font.bold, font.italic, font.name = Pt(para['size']), para['bold'], para['italic'], 'Arial'
        font.color.rgb = _rgb(para['color'])
        rPr = run._r.get_or_add_rPr()
        for attr in ('spc', 'kern'):
            if para[attr] is not None:
                rPr.set(attr, str(para[attr]))
        for tag in ('a:ea', 'a:cs'):
            etree.SubElement(rPr, qn(tag), typeface='Arial')
    if name:
        tb.name = name
    return tb


def _add_icon(slide, key, x, y, w, h):
    return slide.shapes.add_picture(io.BytesIO(_icons(TEMPLATE)[key]), Emu(x), Emu(y), Emu(w), Emu(h))


def _add_slide_number(slide, color, number):
    """자동 갱신되는 슬라이드 번호 필드"""
    tb = _add_text(slide, 8503920, 320040, 800000, 300000, [_para(str(number), 10, color)], anchor='ctr',
                   margins=(91440, 45720, 91440, 45720), name='슬라이드 번호')
    p = tb.text_frame.paragraphs[0]._p
    run = p.find(qn('a:r'))
    fld = etree.Element(qn('a:fld'), id='{F7021451-1387-4CA6-816F-3879F97B5CBC}', type='slidenum')
    fld.append(copy.deepcopy(run.find(qn('a:rPr'))))
    etree.SubElement(fld, qn('a:t')).text = str(number)
    p.replace(run, fld)


def _add_header_tag(slide, tag):
    _add_rect(slide, 457200, 411480, 365760, 365760, DARK)
    _add_icon(slide, 'warning', 548640, 484632, 219456, 219456)
    _add_rect(slide, 822960, 411480, 1600200, 365760, DARK)
    _add_text(slide, 896112, 411480, 1463040, 365760, [_para(tag, 12, YELLOW, bold=True, spc=100, kern=0)], 'ctr')


def _add_footer(slide, color):
    _add_text(slide, 457200, 6510528, 2743200, 274320, [_para('AIGLUE 안전보건팀', 9, color)], 'ctr')


def _add_corner_blocks(slide):
    _add_rect(slide, 7863840, 5760720, 1280160, 1097280, YELLOW)
    _add_rect(slide, 7863840, 5760720, 457200, 457200, DARK)


def _build_cover(prs, date_text):
    s = _new_slide(prs, DARK)
    _add_corner_blocks(s)
    _add_footer(s, GRAY_LIGHT)
    _add_text(s, 548640, 457200, 3657600, 365760, [_para('A I G L U E', 14, WHITE, bold=True, spc=300, kern=0)], 'ctr')
    _add_rect(s, 548640, 1417320, 3017520, 457200, YELLOW)
    _add_text(s, 548640, 1417320, 3017520, 457200,
              [_para('SAFETY FIRST · 정기 안전보건 자료', 11, DARK, bold=True, align='ctr')], 'ctr')
    _add_text(s, 548640, 2377440, 7680960, 1737360,
              [_para('제조현장', 40, WHITE, bold=True, line_pt=46), _para('안전사고 사례', 40, WHITE, bold=True, line_pt=46)])
    _add_text(s, 548640, 4343400, 7315200, 457200, [_para('AIGLUE 안전보건팀', 18, YELLOW, bold=True)])
    _add_text(s, 548640, 5943600, 5486400, 365760, [_para(date_text, 12, GRAY_LIGHT)])
    _add_slide_number(s, GRAY_LIGHT, len(prs.slides))


def _build_toc(prs, items):
    s = _new_slide(prs)
    _add_header_tag(s, 'CONTENTS')
    _add_footer(s, GRAY_MID)
    _add_text(s, 457200, 960120, 5486400, 777240, [_para('목차', 34, DARK, bold=True)])
    for i, (no, title, desc) in enumerate(items):
        y = 2148840 + i * 1051560
        _add_text(s, 457200, y, 914400, 731520, [_para(no, 30, YELLOW, bold=True)])
        _add_text(s, 1508760, y, 6400800, 365760, [_para(title, 18, DARK, bold=True)])
        _add_text(s, 1508760, y + 384048, 6675120, 365760, [_para(desc, 12, GRAY_MID)])
        if i < len(items) - 1:
            _add_rect(s, 457200, y + 868680, 8229600, 10973, RULE)
    _add_slide_number(s, GRAY_MID, len(prs.slides))


def _build_section(prs, number, title, subtitle):
    s = _new_slide(prs, DARK)
    _add_corner_blocks(s)
    _add_footer(s, GRAY_LIGHT)
    _add_text(s, 548640, 1097280, 3657600, 1463040, [_para(number, 80, YELLOW, bold=True)])
    _add_text(s, 548640, 2834640, 7498080, 914400, [_para(title, 32, WHITE, bold=True)])
    _add_text(s, 548640, 3703320, 7315200, 457200, [_para(subtitle, 16, YELLOW)])
    _add_slide_number(s, GRAY_LIGHT, len(prs.slides))


def _fit_in_box(image_bytes, box):
    """사진을 비율을 유지한 채 사진 영역 안에 가운데 정렬로 맞춘다"""
    bx, by, bw, bh = box
    w, h = Image.open(io.BytesIO(image_bytes)).size
    scale = min(bw / w, bh / h)
    pw, ph = int(w * scale), int(h * scale)
    return bx + (bw - pw) // 2, by + (bh - ph) // 2, pw, ph


def _build_case(prs, case: CaseSlide):
    s = _new_slide(prs)
    _add_header_tag(s, '사고 사례')
    _add_footer(s, GRAY_MID)
    _add_text(s, 457200, 914400, 8229600, 548640, [_para(case.title, 25, DARK, bold=True)], name='제목')
    _add_rect(s, 457200, 1691640, 4846320, 4572000, PANEL, shadow=True)
    values = {'일시': [case.date], '장소': [case.place], '소속': [case.team], '내용': list(case.bullets)}
    for icon, label, y in CASE_ROWS:
        _add_icon(s, icon, 731520, y + 27432, 219456, 219456)
        _add_text(s, 1051560, y, 685800, 274320, [_para(label, 12, DARK, bold=True)])
        height = 1417320 if label == '내용' else 274320
        _add_text(s, 1783080, y, 3383280, height, [_para(line, 12, DARK, line_pt=16) for line in values[label]],
                  name=f'{label} 값')
    _add_rect(s, 731520, 4892040, 4297680, 502920, RED)
    _add_text(s, 731520, 4892040, 4297680, 502920, [_para(case.level, 14, WHITE, bold=True, align='ctr')], 'ctr',
              name='재해정도 값')
    fx, fy, fw, fh = PHOTO_FRAME
    _add_rect(s, fx, fy, fw, fh, PHOTO_BG, line=(RULE, 15875, True))
    if case.photo is None:
        _add_icon(s, 'camera', 6766560, 2743200, 640080, 640080)
        _add_text(s, 5669280, 3520440, 2834640, 365760, [_para('사고 현장 사진을 삽입하세요', 11, GRAY_LIGHT, align='ctr')], 'ctr')
    else:
        x, y, w, h = _fit_in_box(case.photo, PHOTO_FRAME)
        s.shapes.add_picture(io.BytesIO(case.photo), Emu(x), Emu(y), Emu(w), Emu(h)).name = '사고 현장 사진'
    _add_text(s, 5486400, 5230368, 3200400, 320040,
              [_para('※ 사진: 현장 촬영본 / 개인정보 노출 금지', 9, GRAY_MID, italic=True)], 'ctr')
    _add_slide_number(s, GRAY_MID, len(prs.slides))


def build_presentation(cases: Sequence[CaseSlide], today: datetime.date | None = None) -> bytes:
    """표지 → 목차(사고사례 항목) → 01 구분 → 사고사례 N장. 완성된 .pptx 파일의 bytes를 반환한다."""
    today = today or datetime.date.today()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    _build_cover(prs, f'{today:%Y. %m. %d}   |   전사 안전보건교육 자료')
    _build_toc(prs, [SECTION_ITEM])
    _build_section(prs, *SECTION_ITEM)
    for case in cases:
        _build_case(prs, case)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ====================================================================================================
# 10. 조회·추출·PPT 로직: 화면과 무관한 순수 함수 (화면 없이도 호출·테스트 가능)
#     조회 → LLM 추출 → 사용자 확인·수정 → 재검증 후 PPT 생성
# ====================================================================================================
TABLE_COLUMNS = ['시트행', 'No', '일시', '사고유형', '재해정도', '소속']
MAX_EXTRACT_ROWS = 12          # extract_rows 한 번에 처리할 최대 케이스 수
EXTRACT_WORKERS = 2            # Upstage 요청 한도(429) 때문에 동시 호출은 2개까지


class ApiError(Exception):
    """클라이언트에 그대로 보여줄 수 있는 오류 (HTTP 상태 코드 + 메시지 + 선택적 상세 목록)"""

    def __init__(self, status: int, message: str, details: list[dict] | None = None) -> None:
        super().__init__(message)
        self.status = status
        self.message = message
        self.details = details or []


@dataclass(frozen=True)
class Dataset:
    source_name: str
    df: pd.DataFrame
    images: Mapping[int, bytes]
    extractor: CaseExtractor | None = None      # None이면 LLM 추출 불가(API 키 없음)


@dataclass(frozen=True)
class QueryRequest:
    date_range: DateRange
    selections: Mapping[str, tuple[str, ...]]
    match: str = 'AND'                        # 변수 사이 결합: 'AND'(기본) 또는 'OR'


@dataclass(frozen=True)
class RowEdit:
    """화면에서 사용자가 확인·수정한 한 케이스의 값"""
    sheet_row: int
    title: str
    place: str
    team: str
    bullets: tuple[str, ...]
    approved: bool = False


def load_dataset(xlsx: ExcelSource = XLSX, extractor: CaseExtractor | None = None, name: str | None = None) -> Dataset:
    source_name = name or (xlsx.name if isinstance(xlsx, Path) else '업로드한 엑셀')
    return Dataset(source_name=source_name, df=build_standard_records(xlsx), images=load_row_images(xlsx),
                   extractor=extractor)


def load_uploaded_dataset(data: bytes, name: str, extractor: CaseExtractor | None) -> Dataset:
    """사용자가 올린 엑셀로 데이터셋을 만든다. 형식이 맞지 않으면 원인을 담은 ExcelFormatError."""
    try:
        _check_excel_limits(data)
        dataset = load_dataset(data, extractor, name=name)
    except ExcelFormatError:
        raise
    except zipfile.BadZipFile as error:
        raise ExcelFormatError('엑셀(.xlsx) 파일이 아니거나 손상된 파일입니다.') from error
    except (KeyError, ValueError, IndexError) as error:
        log.warning('업로드 엑셀 형식 불일치: %s', type(error).__name__)
        raise ExcelFormatError(f"사고관리 엑셀 형식이 아닙니다. '{SHEET}' 시트와 2단 헤더, 이미지가 있는 엑셀을 올려 주세요.") from error
    except Exception as error:                          # 손상·조작된 파일에서 나올 수 있는 그 밖의 예외(상세는 서버 로그에만)
        log.warning('업로드 엑셀을 읽지 못했습니다: %s', type(error).__name__)
        raise ExcelFormatError('엑셀 파일을 읽을 수 없습니다. 사고관리 엑셀(.xlsx)이 맞는지 확인해 주세요.') from error
    if len(dataset.df) > MAX_DATA_ROWS:
        raise ExcelFormatError(f'데이터 행이 너무 많습니다(최대 {MAX_DATA_ROWS:,}행).')
    if not dataset.df['날짜'].notna().any():
        raise ExcelFormatError('사고 데이터가 한 건도 없는 엑셀입니다.')
    return dataset


def build_meta(ds: Dataset) -> dict:
    """화면 초기화용 정보: 일자 범위·날짜별 건수, 변수별 결측 현황과 고유값 목록(콤보 박스 옵션)"""
    first, last = date_bounds(ds.df)
    stats = missing_stats(ds.df).set_index('변수')
    variables = []
    for name in FILTER_VARS:
        table = unique_table(ds.df, name)
        variables.append({
            'name': name,
            'total': int(stats.loc[name, '전체 건수']),
            'missing': int(stats.loc[name, '결측 건수']),
            'missingRatio': float(stats.loc[name, '결측 비율']),
            'options': [{'value': row['값'], 'count': int(row['건수']), 'ratio': float(row['비율'])}
                        for _, row in table.iterrows()],
        })
    return {'sourceName': ds.source_name, 'total': len(ds.df),
            'dateBounds': {'min': first.isoformat(), 'max': last.isoformat()},
            'casesPerDay': cases_per_day(ds.df), 'variables': variables,
            'llmAvailable': ds.extractor is not None,
            'limits': {'title': LIMITS.title, 'place': LIMITS.place, 'team': LIMITS.team,
                       'bodyLine': LIMITS.body_line, 'bodyMaxLines': LIMITS.body_max_lines}}


def parse_query(payload: object, ds: Dataset) -> QueryRequest:
    """클라이언트 입력을 검증해 QueryRequest로 만든다. 잘못된 입력은 ApiError(400)."""
    if not isinstance(payload, dict):
        raise ApiError(400, '요청 형식이 올바르지 않습니다.')
    first, last = date_bounds(ds.df)
    try:
        date_range = normalize_date_range(payload.get('dateFrom') or first, payload.get('dateTo') or last)
    except (ValueError, TypeError) as error:
        raise ApiError(400, f'일자가 올바르지 않습니다: {error}') from error

    raw = payload.get('selections') or {}
    if not isinstance(raw, dict):
        raise ApiError(400, '선택 조건 형식이 올바르지 않습니다.')
    selections: dict[str, tuple[str, ...]] = {}
    for name, value in raw.items():
        if name not in FILTER_VARS:
            raise ApiError(400, f'선택할 수 없는 변수입니다: {name}')
        if value in (None, ''):
            continue
        if not isinstance(value, str):
            raise ApiError(400, f'{name} 값은 문자열이어야 합니다.')
        if value not in set(ds.df[name].unique()):
            raise ApiError(400, f'{name}에 없는 값입니다: {value}')
        selections[name] = (value,)
    match = payload.get('match') or 'AND'
    if match not in MATCH_MODES:
        raise ApiError(400, f'조건 결합은 AND 또는 OR여야 합니다: {match}')
    return QueryRequest(date_range=date_range, selections=selections, match=match)


def _query_df(ds: Dataset, request: QueryRequest) -> pd.DataFrame:
    return filter_cases(ds.df, request.selections, request.date_range, request.match)


def run_query(ds: Dataset, request: QueryRequest) -> dict:
    """조회: 건수·표 행. 0건이면 원인 힌트(조건을 하나씩 뺐을 때의 건수)를 함께 준다."""
    result = _query_df(ds, request)
    rows = [{'시트행': int(r['시트행']), 'No': int(r['No']), '일시': r['일시'].strftime('%Y-%m-%d %H:%M'),
             '사고유형': str(r['사고유형']), '재해정도': level_code(r), '소속': default_team(r)}
            for r in result.to_dict('records')]
    hints = diagnose_empty(ds.df, request.selections, request.date_range, request.match) if result.empty else {}
    return {'count': len(result), 'columns': TABLE_COLUMNS, 'rows': rows, 'hints': hints}


def _records_by_sheet_row(ds: Dataset, sheet_rows: Sequence[int]) -> list[dict]:
    known = ds.df.set_index('시트행')
    missing = [row for row in sheet_rows if row not in known.index]
    if missing:
        raise ApiError(400, f'없는 시트행입니다: {missing}')
    return [{**known.loc[row].to_dict(), '시트행': row} for row in sheet_rows]


def _extract_one(extractor: CaseExtractor, record: dict) -> dict:
    sheet_row = int(record['시트행'])
    try:
        extraction = extract_case(extractor, str(record['사고내용']))
    except (UpstageError, ApiKeyMissing) as error:
        return {'sheetRow': sheet_row, 'error': getattr(error, 'message', str(error))}
    except Exception:
        log.exception('추출 중 오류 (시트행 %s)', sheet_row)
        return {'sheetRow': sheet_row, 'error': '추출 중 오류가 발생했습니다.'}
    rule = rule_based_place(str(record['사고내용']))
    place, differs = default_place(rule, extraction.place_llm)
    return {'sheetRow': sheet_row, 'title': extraction.title, 'placeLlm': extraction.place_llm, 'placeRule': rule,
            'placeDefault': place, 'placeDiffers': differs, 'bullets': list(extraction.bullets),
            'model': extraction.model, 'tokens': extraction.tokens, 'warnings': list(extraction.warnings)}


def extract_rows(ds: Dataset, payload: object) -> list[dict]:
    """조회된 케이스의 제목·장소·내용을 LLM(+규칙)으로 추출한다. 실패한 케이스는 error 항목으로 돌려준다."""
    extractor = ds.extractor
    if extractor is None:
        raise ApiError(503, 'LLM을 쓸 수 없습니다. UPSTAGE_API_KEY를 .env 파일 또는 Streamlit Secrets에 넣고 앱을 다시 실행해 주세요.')
    rows = payload.get('sheetRows') if isinstance(payload, dict) else None
    if not (isinstance(rows, list) and rows and all(isinstance(r, int) and not isinstance(r, bool) for r in rows)):
        raise ApiError(400, 'sheetRows는 정수 목록이어야 합니다.')
    if len(rows) > MAX_EXTRACT_ROWS:
        raise ApiError(400, f'한 번에 {MAX_EXTRACT_ROWS}건까지만 추출할 수 있습니다.')
    records = _records_by_sheet_row(ds, rows)
    with ThreadPoolExecutor(EXTRACT_WORKERS) as pool:
        return list(pool.map(lambda record: _extract_one(extractor, record), records))


def _parse_edits(payload: object) -> dict[int, RowEdit]:
    if not isinstance(payload, list):
        raise ApiError(400, 'edits는 목록이어야 합니다.')
    edits: dict[int, RowEdit] = {}
    for item in payload:
        try:
            edits[int(item['sheetRow'])] = RowEdit(
                sheet_row=int(item['sheetRow']), title=str(item['title']).strip(), place=str(item['place']).strip(),
                team=str(item['team']).strip(), bullets=tuple(str(b) for b in item['bullets']),
                approved=bool(item.get('approved', False)))
        except (KeyError, TypeError, ValueError) as error:
            raise ApiError(400, '수정 내용의 형식이 올바르지 않습니다.') from error
    return edits


def field_problems(title: str, place: str, team: str, bullets: Sequence[str]) -> list[str]:
    """칸별 비어 있음·넘침 문제 목록 (Streamlit 화면과 PPT 생성 단계가 같은 검사를 쓴다)"""
    found: list[str] = []
    fields = (('제목', title, title_issue(title) if title else None),
              ('장소', place, single_line_issue('장소', place, LIMITS.place) if place else None),
              ('소속', team, single_line_issue('소속', team, LIMITS.team) if team else None))
    for label, value, issue in fields:
        if not value:
            found.append(f'{label}이(가) 비어 있습니다')
        elif issue:
            found.append(issue)
    cleaned = [b.strip().lstrip('-·•').strip() for b in bullets if b.strip()]
    if not cleaned:
        found.append('내용이 비어 있습니다')
    elif (issue := body_issue(cleaned)):
        found.append(issue)
    return found


def _edit_problems(ds: Dataset, record: dict, edit: RowEdit) -> list[str]:
    """한 케이스의 넘침·누락·미승인 문제 목록 (PPT 칸을 넘치거나 비면 만들지 않는다)"""
    found = field_problems(edit.title, edit.place, edit.team, edit.bullets)
    extractor = ds.extractor
    if extractor is not None and not edit.approved:
        try:
            extraction = extract_case(extractor, str(record['사고내용']))
            _, differs = default_place(rule_based_place(str(record['사고내용'])), extraction.place_llm)
        except (UpstageError, ApiKeyMissing):
            differs = False
        if differs:
            found.append('장소가 규칙 추출값과 LLM 추출값이 달라 승인이 필요합니다')
    return found


def build_ppt(ds: Dataset, request: QueryRequest, edits_payload: object) -> tuple[bytes, dict[str, int]]:
    """조회 결과 + 사용자가 확인·수정한 값으로 PPT 생성. 조건을 다시 조회하므로 표와 항상 같은 케이스가 들어간다."""
    result = _query_df(ds, request)
    if result.empty:
        raise ApiError(409, '조회된 데이터가 없어 PPT를 만들 수 없습니다.')
    edits = _parse_edits(edits_payload)
    problems, cases = [], []
    for record in result.to_dict('records'):
        row = int(record['시트행'])
        edit = edits.get(row)
        if edit is None:
            problems.append({'sheetRow': row, 'no': int(record['No']), 'message': '확인·수정 내용이 없습니다'})
            continue
        problems += [{'sheetRow': row, 'no': int(record['No']), 'message': m} for m in _edit_problems(ds, record, edit)]
        cases.append(to_case_slide(record, title=edit.title, place=edit.place, team=edit.team,
                                   bullets=edit.bullets, photo=ds.images.get(row)))
    if problems:
        raise ApiError(422, f'{len(problems)}건의 문제를 해결해야 PPT를 만들 수 없습니다.', problems)
    return build_presentation(cases), {'count': len(cases)}


# ====================================================================================================
# 11. Streamlit 화면 (조회 → LLM 추출 → 확인·수정 표 → 확인 후 PPT)
#     실행: streamlit run app.py  (또는 streamlit run all_function_code.py)
#     표 편집·검증·PPT 생성은 10번의 함수(run_query, extract_rows, field_problems, build_ppt)를 그대로 쓴다.
# ====================================================================================================
BULLET_SEPARATOR = ' | '             # 표의 '내용' 칸 한 셀에 개조식 항목을 이어 쓰는 구분자
NO_SELECTION = '(전체)'
EXTRACT_CHUNK = 4                    # 진행률을 보여주며 4건씩 추출 요청
LOCKED_COLUMNS = ['시트행', 'No', '일시', '재해정도', '규칙≠LLM', '주의']    # 사용자가 고칠 수 없는 칸


def split_bullets(text) -> list[str]:
    """표의 '내용' 셀 → 개조식 항목 목록 (구분자 '|', 앞의 '- '는 제거)"""
    items = (item.strip().lstrip('-·•').strip() for item in str(text or '').split('|'))
    return [item for item in items if item]


def build_review_table(rows: Sequence[dict], extracted: Mapping[int, dict]) -> pd.DataFrame:
    """조회 행(run_query) + 추출 결과 → 확인·수정 표의 처음 값. 추출에 실패한 행은 칸을 비워 사용자가 직접 쓰게 한다."""
    records = []
    for row in rows:
        info = extracted.get(row['시트행'], {})
        failed = 'error' in info or not info
        note = f"추출 실패: {info.get('error', '아직 추출되지 않음')}" if failed else '; '.join(info.get('warnings', []))
        records.append({'시트행': row['시트행'], 'No': row['No'], '일시': row['일시'], '재해정도': row['재해정도'],
                        '소속': row['소속'], '제목': '' if failed else info['title'],
                        '장소': '' if failed else info['placeDefault'],
                        '내용': '' if failed else BULLET_SEPARATOR.join(info['bullets']),
                        '규칙≠LLM': bool(info.get('placeDiffers', False)), '장소 승인': False, '주의': note})
    return pd.DataFrame(records)


def edits_from_table(table: pd.DataFrame, extracted: Mapping[int, dict]) -> list[dict]:
    """사용자가 수정한 표 → build_ppt가 받는 edits 형식. 장소를 직접 고치면 승인한 것으로 본다."""
    edits = []
    for row in table.to_dict('records'):
        sheet_row = int(row['시트행'])
        place = str(row['장소'] or '').strip()
        default_place_text = extracted.get(sheet_row, {}).get('placeDefault', place)
        edits.append({'sheetRow': sheet_row, 'title': str(row['제목'] or '').strip(), 'place': place,
                      'team': str(row['소속'] or '').strip(), 'bullets': split_bullets(row['내용']),
                      'approved': bool(row['장소 승인']) or place != default_place_text})
    return edits


def review_problems(edits: Sequence[dict], extracted: Mapping[int, dict]) -> dict[int, list[str]]:
    """시트행 → 해결해야 할 문제 목록 (비어 있음·넘침·규칙/LLM 장소 미승인). 없는 행은 키가 없다."""
    found: dict[int, list[str]] = {}
    for edit in edits:
        issues = field_problems(edit['title'], edit['place'], edit['team'], edit['bullets'])
        if extracted.get(edit['sheetRow'], {}).get('placeDiffers') and not edit['approved']:
            issues.append('장소가 규칙 추출값과 LLM 추출값이 달라 승인이 필요합니다')
        if issues:
            found[edit['sheetRow']] = issues
    return found


def apply_row_fix(table: pd.DataFrame, sheet_row: int, *, title: str, place: str, team: str, bullets_text: str,
                  approved: bool) -> pd.DataFrame:
    """문제 행 수정 창에서 고친 값을 표의 그 행에 반영한 새 표를 돌려준다 (원본 표는 바꾸지 않음).
    bullets_text는 한 줄에 항목 하나."""
    fixed = table.copy()
    index = fixed.index[fixed['시트행'] == sheet_row]
    if len(index) != 1:
        raise KeyError(f'표에 없는 시트행입니다: {sheet_row}')
    bullets = split_bullets(bullets_text.replace('\n', '|'))
    for column, value in (('제목', title.strip()), ('장소', place.strip()), ('소속', team.strip()),
                          ('내용', BULLET_SEPARATOR.join(bullets)), ('장소 승인', bool(approved))):
        fixed.loc[index[0], column] = value
    return fixed


def merge_retry(table: pd.DataFrame, fresh: pd.DataFrame, sheet_rows: Sequence[int]) -> pd.DataFrame:
    """다시 추출한 행(sheet_rows)만 새 값으로 바꾸고, 사용자가 고친 다른 행은 그대로 둔 새 표"""
    merged = table.copy()
    fresh_by_row = fresh.set_index('시트행')
    for row in sheet_rows:
        index = merged.index[merged['시트행'] == row][0]
        for column in ('제목', '장소', '내용', '규칙≠LLM', '주의'):
            merged.loc[index, column] = fresh_by_row.loc[row, column]
    return merged


def downloads_folder() -> Path:
    """이 PC의 다운로드 폴더 (Windows는 이동·변경된 위치까지 레지스트리에서 찾는다)"""
    if os.name == 'nt':
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                                r'Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders') as key:
                value, _ = winreg.QueryValueEx(key, '{374DE290-123F-4565-9164-39C4925E467B}')
            return Path(os.path.expandvars(value))
        except OSError:
            pass
    return Path.home() / 'Downloads'


def open_downloads_folder(delay: float = 1.5, opener: Callable[[str], object] | None = None) -> None:
    """다운로드 버튼을 누른 직후 브라우저가 파일을 저장할 시간(delay초)을 준 뒤 다운로드 폴더를 연다.
    이 앱을 실행 중인 PC에서만 의미가 있고, 폴더를 열지 못해도 화면 동작에는 영향이 없다."""
    folder = downloads_folder()
    if not folder.is_dir():
        log.warning('다운로드 폴더를 찾을 수 없습니다: %s', folder)
        return
    open_folder = opener or getattr(os, 'startfile', None) or (lambda path: webbrowser.open(Path(path).as_uri()))
    timer = threading.Timer(delay, open_folder, args=(str(folder),))
    timer.daemon = True
    timer.start()


def _is_desktop_os() -> bool:
    return os.name == 'nt' or sys.platform == 'darwin'


def can_open_downloads_folder() -> bool:
    """이 앱이 사용자의 PC에서 실행 중이라 다운로드 폴더를 열 수 있는지.
    리눅스 서버(클라우드)에서는 폴더가 우연히 있어도 서버에서 파일 탐색기를 여는 일이 없도록 항상 False."""
    return _is_desktop_os() and downloads_folder().is_dir()


def _load_default_dataset() -> Dataset:
    return load_dataset(XLSX, build_extractor())


def _st_extractor() -> CaseExtractor | None:
    """모든 접속이 함께 쓰는 LLM 추출기 (키·캐시 파일 한 벌)"""
    return st.cache_resource(show_spinner=False)(build_extractor)()


def _st_dataset() -> Dataset:
    """'엑셀 파일 열기'에서 사용자가 고른 엑셀로 데이터셋을 만든다.
    고르지 않았을 때는 폴더에 기본 엑셀이 있으면 그것을 쓰고(내 PC), 없으면(클라우드) 고르라고 안내하고 멈춘다."""
    uploaded = st.file_uploader('엑셀 파일 열기 (.xlsx)', type=['xlsx'], key='excel_upload',
                                help='사고관리 엑셀을 파일 선택 버튼으로 고르거나 이 칸에 끌어다 놓습니다. '
                                     '고른 파일은 이 접속(세션) 동안만 서버에서 쓰이고 저장소에는 저장되지 않습니다.')
    if uploaded is None:
        if st.session_state.pop('upload_dataset', None) is not None:
            _st_reset_source_state()
        if XLSX.exists():
            st.caption(f'기본 엑셀({XLSX.name})을 쓰고 있습니다. 다른 엑셀을 쓰려면 위 칸에서 파일을 고르세요.')
            return st.cache_resource(show_spinner='엑셀을 읽는 중입니다...')(_load_default_dataset)()
        st.info('위의 "엑셀 파일 열기"에서 사고관리 엑셀(.xlsx)을 선택(업로드)해 주세요.')
        st.stop()
    signature = uploaded.file_id
    cached = st.session_state.get('upload_dataset')
    if cached is None or cached[0] != signature:
        try:
            with st.spinner('엑셀을 읽는 중입니다...'):
                dataset = load_uploaded_dataset(uploaded.getvalue(), uploaded.name, _st_extractor())
        except ExcelFormatError as error:
            st.error(str(error))
            if XLSX.exists():
                st.caption('이 파일은 쓸 수 없어 화면을 멈췄습니다. 위 칸에서 파일을 지우면 기본 엑셀로 돌아갑니다.')
            st.stop()
        _st_reset_source_state()
        st.session_state['upload_dataset'] = (signature, dataset)
    return st.session_state['upload_dataset'][1]


def _st_reset_results() -> None:
    for key in ('q_payload', 'q_result', 'q_extracted', 'q_table', 'q_ppt', 'q_declined'):
        st.session_state.pop(key, None)


def _st_reset_source_state() -> None:
    """데이터 소스(기본 엑셀·고른 엑셀)가 바뀌면 이전 데이터 기준의 조회 결과와 조건 위젯 값(일자·선택 항목·표 편집)을 모두 버린다.
    남겨 두면 새 엑셀의 일자 범위·선택지를 벗어난 값이 위젯에 남아 화면 오류가 날 수 있다."""
    _st_reset_results()
    for key in list(st.session_state.keys()):
        if key in ('date_range', 'match') or str(key).startswith(('sel_', 'editor_')):
            st.session_state.pop(key, None)


def _st_filters(meta: dict) -> dict | None:
    """일자 범위·조건 결합(AND/OR)·변수별 콤보 박스를 그리고 조회 요청(payload)을 돌려준다. 일자가 덜 골라졌으면 None."""
    first = datetime.date.fromisoformat(meta['dateBounds']['min'])
    last = datetime.date.fromisoformat(meta['dateBounds']['max'])
    picked = st.date_input('사고 일자 (시작일 ~ 종료일)', value=(first, last), min_value=first, max_value=last, key='date_range')
    match = st.radio('조건 검색 방식 (사고유형·재해정도·사고성여부 사이)', MATCH_MODES, index=0, horizontal=True, key='match',
                     format_func=lambda m: 'AND — 선택한 조건을 모두 만족' if m == 'AND' else 'OR — 선택한 조건 중 하나라도 만족')
    st.caption(f"'{NO_SELECTION}'를 고른 변수는 조건에서 무시됩니다. 일자 범위는 항상 함께 적용됩니다.")
    selections: dict[str, str] = {}
    for column, variable in zip(st.columns(len(meta['variables']), gap='large'), meta['variables']):
        with column:
            counts = {option['value']: option['count'] for option in variable['options']}
            choice = st.selectbox(variable['name'], [NO_SELECTION, *counts], key=f"sel_{variable['name']}",
                                  format_func=lambda v, c=counts: v if v == NO_SELECTION else f'{v} ({c[v]}건)')
            st.caption(f"결측 {variable['missing']:,}건 / 전체 {variable['total']:,}건 ({variable['missingRatio']:.1%})")
            if choice != NO_SELECTION:
                selections[variable['name']] = choice
    if not (isinstance(picked, (tuple, list)) and len(picked) == 2):
        st.info('종료일까지 골라 주세요. 하루만 조회하려면 같은 날짜를 두 번 선택합니다.')
        return None
    return {'dateFrom': picked[0].isoformat(), 'dateTo': picked[1].isoformat(), 'selections': selections, 'match': match}


def describe_conditions(payload: Mapping) -> str:
    """조회 조건을 한 줄로: '사고유형=끼임 AND 재해정도=FAC' (선택 없으면 '조건 없음(전체)')"""
    parts = [f'{name}={value}' for name, value in payload['selections'].items()]
    return f" {payload.get('match', 'AND')} ".join(parts) if parts else '조건 없음(전체)'


def _st_extract_all(ds: Dataset, sheet_rows: Sequence[int]) -> dict[int, dict]:
    """LLM 추출을 진행률과 함께 실행한다 (한 번에 EXTRACT_CHUNK건)"""
    results: dict[int, dict] = {}
    bar = st.progress(0.0, text='사고내용에서 제목·장소·내용을 추출하는 중입니다...')
    for start in range(0, len(sheet_rows), EXTRACT_CHUNK):
        chunk = list(sheet_rows[start:start + EXTRACT_CHUNK])
        results.update({item['sheetRow']: item for item in extract_rows(ds, {'sheetRows': chunk})})
        done = start + len(chunk)
        bar.progress(done / len(sheet_rows), text=f'추출 {done}/{len(sheet_rows)}건')
    bar.empty()
    return results


def _st_run_query(ds: Dataset, payload: dict) -> None:
    """'선택한 조건으로 조회': 조회 → (결과가 있고 LLM을 쓸 수 있으면) 곧바로 추출까지 실행한다."""
    try:
        result = run_query(ds, parse_query(payload, ds))
    except ApiError as error:
        st.error(error.message)
        return
    version = st.session_state.get('q_version', 0) + 1
    _st_reset_results()
    extracted: dict[int, dict] = {}
    if result['count'] and ds.extractor is not None:
        extracted = _st_extract_all(ds, [row['시트행'] for row in result['rows']])
    st.session_state.update(q_payload=payload, q_result=result, q_extracted=extracted, q_version=version,
                            q_table=build_review_table(result['rows'], extracted))


def _st_retry_failed(ds: Dataset, failed: Sequence[int], current: pd.DataFrame) -> None:
    """추출에 실패한 행만 다시 추출한다. 사용자가 이미 고친 다른 행은 그대로 둔다."""
    state = st.session_state
    state['q_extracted'].update(_st_extract_all(ds, list(failed)))
    fresh = build_review_table(state['q_result']['rows'], state['q_extracted'])
    state['q_table'] = merge_retry(current, fresh, failed)
    state['q_version'] += 1                      # 표를 새 값으로 다시 그린다
    st.rerun()


def _st_editor(table: pd.DataFrame, version: int) -> pd.DataFrame:
    st.caption(f"칸 한도(한글 기준): 제목 {LIMITS.title:.0f} (앞의 '[사고사례] ' 포함) · 장소 {LIMITS.place:.0f} · "
               f"소속 {LIMITS.team:.0f} · 내용 줄당 {LIMITS.body_line:.0f}자 × 최대 {LIMITS.body_max_lines}줄. "
               "내용은 항목을 '|'로 구분해 씁니다.")
    config = {'제목': st.column_config.TextColumn('제목 (LLM 생성·수정 가능)', width='medium'),
              '소속': st.column_config.TextColumn('소속 (관리/책임부서)', width='medium'),
              '장소': st.column_config.TextColumn('장소 (규칙 / LLM)', width='medium'),
              '내용': st.column_config.TextColumn('내용 (개조식, | 로 구분)', width='large'),
              '규칙≠LLM': st.column_config.CheckboxColumn('규칙≠LLM', help='장소를 규칙과 LLM이 다르게 찾음: 확인 후 승인'),
              '장소 승인': st.column_config.CheckboxColumn(
                  '장소 승인', help='규칙≠LLM일 때 이대로 써도 되면 체크 (장소를 직접 고치면 자동 승인)'),
              '재해정도': st.column_config.TextColumn('재해정도 (원문 코드)')}
    return st.data_editor(table, key=f'editor_{version}', hide_index=True, width='stretch', disabled=LOCKED_COLUMNS,
                          column_config=config, num_rows='fixed')


def _st_problem_picker(problems: Mapping[int, list[str]], numbers: Mapping[int, int], version: int) -> int | None:
    """문제 표를 보여주고, 사용자가 고른(클릭한) 행의 시트행을 돌려준다. 고르지 않았으면 None."""
    frame = pd.DataFrame([{'시트행': row, 'No': numbers[row], '문제': ' / '.join(found)} for row, found in problems.items()])
    st.caption('문제 표의 행을 클릭하면 아래에 그 행의 수정 칸이 열립니다. (Streamlit 표는 더블클릭 이벤트가 없어 한 번 클릭으로 대신합니다.)')
    event = st.dataframe(frame, hide_index=True, width='stretch', on_select='rerun', selection_mode='single-row',
                         key=f'problem_pick_{version}')
    picked = event.selection.rows if event and event.selection else []
    return int(frame.iloc[picked[0]]['시트행']) if picked else None


def _st_apply_fix(table: pd.DataFrame, sheet_row: int, key: str) -> None:
    """'적용' 버튼 콜백: 수정 창의 값을 표에 반영하고, 표·문제 표를 새 값으로 다시 그리도록 버전을 올린다."""
    state = st.session_state
    state['q_table'] = apply_row_fix(table, sheet_row, title=state[f'{key}_title'], place=state[f'{key}_place'],
                                     team=state[f'{key}_team'], bullets_text=state[f'{key}_bullets'],
                                     approved=bool(state.get(f'{key}_approved', False)))
    state['q_version'] += 1


def _st_fix_panel(ds: Dataset, table: pd.DataFrame, sheet_row: int, extracted: Mapping[int, dict], version: int) -> None:
    """고른 문제 행을 한 곳에서 고친다: 사고내용 원문·넘침 표시를 보며 수정 → '적용'하면 표에 반영된다."""
    row = table[table['시트행'] == sheet_row].iloc[0]
    key = f'fix_{version}_{sheet_row}'
    st.subheader(f'시트행 {sheet_row} (No {row["No"]}) 수정')
    source = ds.df.loc[ds.df['시트행'] == sheet_row, '사고내용'].iloc[0]
    with st.expander('사고내용 원문', expanded=False):
        st.write(source or '(사고내용 없음)')
    title = st.text_input("제목 (앞에 '[사고사례] '가 자동으로 붙습니다)", row['제목'], key=f'{key}_title')
    place = st.text_input('장소', row['장소'], key=f'{key}_place')
    team = st.text_input('소속', row['소속'], key=f'{key}_team')
    bullets_text = st.text_area('내용 (한 줄에 개조식 항목 하나)', '\n'.join(split_bullets(row['내용'])), height=150, key=f'{key}_bullets')
    differs = bool(extracted.get(sheet_row, {}).get('placeDiffers'))
    approved = st.checkbox('장소 승인 (규칙과 LLM 장소가 달라 확인이 필요한 행)', value=bool(row['장소 승인']),
                           key=f'{key}_approved') if differs else False
    issues = field_problems(title.strip(), place.strip(), team.strip(), split_bullets(bullets_text.replace('\n', '|')))
    if differs and not approved and place.strip() == extracted[sheet_row]['placeDefault']:
        issues.append('장소 승인이 필요합니다 (체크하거나 장소를 직접 고치세요)')
    (st.warning if issues else st.success)(' / '.join(issues) if issues else '문제가 없습니다. 적용을 누르세요.')
    st.button('적용', type='primary', key=f'{key}_apply', on_click=_st_apply_fix, args=(table, sheet_row, key))


def _st_build_ppt(ds: Dataset, payload: dict, edits: list[dict], signature: str) -> None:
    """build_ppt와 같은 검증을 거쳐 PPT를 만들고 결과(또는 문제 목록)를 화면 상태에 기록한다."""
    try:
        data, info = build_ppt(ds, parse_query(payload, ds), edits)
    except ApiError as error:
        st.error(error.message)
        if error.details:
            st.dataframe(pd.DataFrame(error.details), hide_index=True, width='stretch')
        return
    st.session_state['q_ppt'] = {'data': data, 'count': info['count'], 'signature': signature,
                                 'name': f'사고사례_{datetime.datetime.now():%Y%m%d_%H%M%S}.pptx'}


def _st_ppt_step(ds: Dataset, payload: dict, edits: list[dict], blocked: bool) -> None:
    """확인 질문 → 예/아니요 → PPT 생성 → 내려받기(누르면 다운로드 폴더가 열린다)"""
    signature = json.dumps([payload, edits], ensure_ascii=False, sort_keys=True)
    st.subheader('PPT 생성')
    st.write(f'조회된 **{len(edits):,}건**으로 PPT를 만들까요?')
    yes, no, _ = st.columns([1, 1, 4])
    if yes.button('예, PPT 만들기', type='primary', disabled=blocked):
        st.session_state.pop('q_declined', None)
        _st_build_ppt(ds, payload, edits, signature)
    if no.button('아니요'):
        st.session_state.pop('q_ppt', None)
        st.session_state['q_declined'] = True
    if st.session_state.get('q_declined'):
        st.info('PPT를 만들지 않았습니다. 내용을 고치거나 조건을 바꿔 다시 조회할 수 있습니다.')
    ppt = st.session_state.get('q_ppt')
    if ppt and ppt['signature'] != signature:
        st.info('표의 내용이 바뀌어 이전에 만든 PPT는 최신이 아닙니다. 다시 만들어 주세요.')
    elif ppt:
        st.success(f"PPT를 만들었습니다. (사고사례 {ppt['count']}장)")
        can_open = can_open_downloads_folder()
        st.download_button('PPT 내려받기', data=ppt['data'], file_name=ppt['name'], mime=PPTX_MIME,
                           on_click=open_downloads_folder if can_open else None)
        if can_open:
            st.caption(f'내려받기를 누르면 잠시 뒤 다운로드 폴더({downloads_folder()})가 열립니다.')


def _st_review_step(ds: Dataset, payload: dict) -> None:
    state = st.session_state
    result, extracted = state['q_result'], state['q_extracted']
    if result['count'] == 0:
        st.warning('조회된 데이터가 없습니다. 선택한 조건에 해당하는 사고 케이스가 엑셀에 없습니다.')
        if result['hints']:
            st.caption('조건을 하나씩 뺐을 때의 건수: ' + ', '.join(f'{k} {n}건' for k, n in result['hints'].items()))
        return
    st.success(f"선택한 조건으로 엑셀에서 {result['count']:,}건이 조회되었습니다. (행번호 순)")
    st.caption(f'적용된 조건: 일자 {payload["dateFrom"]} ~ {payload["dateTo"]} · {describe_conditions(payload)}')
    if ds.extractor is None:
        st.error('LLM을 쓸 수 없습니다. UPSTAGE_API_KEY를 .env 파일 또는 Streamlit Secrets에 넣고 앱을 다시 실행해 주세요.')
        return
    st.subheader('조회 결과 확인·수정')
    edited = _st_editor(state['q_table'], state['q_version'])
    edits = edits_from_table(edited, extracted)
    problems = review_problems(edits, extracted)
    failed = [row for row, info in extracted.items() if 'error' in info]
    if failed:
        st.warning(f'{len(failed)}건은 추출에 실패했습니다. 다시 추출하거나 표에 직접 쓰세요.')
        if st.button('실패한 건 다시 추출'):
            _st_retry_failed(ds, failed, edited)
    approvals = sum(1 for e in edits if extracted.get(e['sheetRow'], {}).get('placeDiffers') and not e['approved'])
    col_a, col_b, col_c = st.columns(3)
    col_a.metric('조회 건수', f'{len(edits):,}')
    col_b.metric('해결할 문제가 있는 건', f'{len(problems):,}')
    col_c.metric('장소 승인이 필요한 건', f'{approvals:,}')
    if problems:
        st.error('아래 문제를 모두 해결해야 PPT를 만들 수 있습니다.')
        picked = _st_problem_picker(problems, {r['시트행']: r['No'] for r in result['rows']}, state['q_version'])
        if picked is not None:
            _st_fix_panel(ds, edited, picked, extracted, state['q_version'])
    _st_ppt_step(ds, payload, edits, blocked=bool(problems))


def render_streamlit_app() -> None:
    """Streamlit 화면 전체: 조건 선택 → '선택한 조건으로 조회'(추출 포함) → 확인·수정 → PPT"""
    st.set_page_config(page_title='사고사례 PPT 생성', layout='wide')
    st.title('사고사례 PPT 생성')
    ds = _st_dataset()
    meta = build_meta(ds)
    st.caption(f"데이터: {meta['sourceName']} · 전체 {meta['total']:,}건 · 결측은 엑셀에서 비어 있는 값입니다.")
    payload = _st_filters(meta)
    if payload is not None and st.button('선택한 조건으로 조회', type='primary'):
        _st_run_query(ds, payload)
    if 'q_result' not in st.session_state:
        return
    st.divider()
    if st.session_state['q_payload'] != payload:
        st.info('선택 조건이 바뀌었습니다. 다시 조회해 주세요.')
        return
    _st_review_step(ds, payload)


def _running_in_streamlit() -> bool:
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx
    except ImportError:
        return False
    return get_script_run_ctx() is not None


if __name__ == '__main__' and _running_in_streamlit():      # streamlit run all_function_code.py
    render_streamlit_app()
